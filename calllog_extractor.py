import subprocess
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

def extract_call_logs():
    try:
        adb_path = r"C:\platform-tools\adb.exe"

        result = subprocess.run(
            [adb_path, "shell", "content", "query", "--uri", "content://call_log/calls"],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True
        )

        if result.returncode != 0:
            print("❌ Error fetching call logs:", result.stderr)
            return

        rows = []
        for line in result.stdout.splitlines():
            if "number=" in line and "type=" in line and "duration=" in line and "date=" in line:
                number = line.split("number=")[1].split(",")[0].strip()
                call_type_code = line.split("type=")[1].split(",")[0].strip()
                duration = line.split("duration=")[1].split(",")[0].strip()
                date_raw = line.split("date=")[1].split(",")[0].strip()

                call_type_map = {"1": "INCOMING", "2": "OUTGOING", "3": "MISSED"}
                call_type = call_type_map.get(call_type_code, call_type_code)

                try:
                    date_time = datetime.fromtimestamp(int(date_raw) / 1000)
                except:
                    date_time = None

                rows.append([number, call_type, int(duration), date_time])

        df = pd.DataFrame(rows, columns=["Number", "Type", "Duration (sec)", "Date & Time"])

        # Top 10 numbers by total duration
        top10 = df.groupby("Number")["Duration (sec)"].sum().sort_values(ascending=False).head(10).reset_index()

        # Save both DataFrames into Excel
        output_file = "../data/call_logs.xlsx"
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Call Logs", index=False)
            top10.to_excel(writer, sheet_name="Top 10 Numbers", index=False)

        # Add bar chart into "Top 10 Numbers" sheet
        wb = load_workbook(output_file)
        ws = wb["Top 10 Numbers"]

        chart = BarChart()
        chart.title = "Top 10 Numbers (by Total Duration)"
        chart.x_axis.title = "Phone Number"
        chart.y_axis.title = "Total Duration (sec)"

        data = Reference(ws, min_col=2, min_row=1, max_row=11)  # Duration column
        cats = Reference(ws, min_col=1, min_row=2, max_row=11)  # Numbers
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)

        ws.add_chart(chart, "E5")  # place chart near the table

        wb.save(output_file)

        print(f"✅ Call Logs extracted and saved with chart to {output_file}")

    except Exception as e:
        print("❌ Exception occurred:", e)


extract_call_logs()
